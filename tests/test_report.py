"""The Excel report: sheet layout, totals that survive the pivot, and formatting.

Workbooks are written to tmp_path — the repo's output/ folder is left alone.
"""

from __future__ import annotations

import openpyxl
import pandas as pd
import pytest

from conftest import EXPECTED, TOTAL_TRANSACTIONS
from merge_statements import Statement, build_report

TRANSACTION_COLUMNS = ["Month", "Date", "Description", "Category", "Amount", "Source File"]
SUMMARY_COLUMNS = [
    "Month",
    "Transactions",
    "Inflows",
    "Outflows",
    "Net",
    "Stated Net",
    "Reconciled",
    "Notes",
]


@pytest.fixture(scope="module")
def report(tmp_path_factory, parsed_statements):
    """One workbook built from all 12 fixtures, reused across the module."""
    path = tmp_path_factory.mktemp("report") / "report.xlsx"
    summary = build_report(parsed_statements, path)
    return path, summary


def sheet(path, name) -> pd.DataFrame:
    return pd.read_excel(path, sheet_name=name)


def test_writes_the_workbook_and_returns_the_summary(report):
    path, summary = report
    assert path.exists()
    assert list(summary.columns) == SUMMARY_COLUMNS
    assert len(summary) == len(EXPECTED)


def test_creates_missing_parent_directories(tmp_path, parsed_statements):
    nested = tmp_path / "does" / "not" / "exist" / "report.xlsx"
    build_report(parsed_statements[:2], nested)
    assert nested.exists()


def test_has_exactly_the_three_documented_sheets(report):
    path, _ = report
    workbook = openpyxl.load_workbook(path)
    assert workbook.sheetnames == ["Summary", "By Category", "Transactions"]


# --- Transactions ---------------------------------------------------------


def test_transactions_sheet_holds_every_row(report):
    path, _ = report
    tx = sheet(path, "Transactions")
    assert list(tx.columns) == TRANSACTION_COLUMNS
    assert len(tx) == TOTAL_TRANSACTIONS


def test_transactions_are_sorted_by_date(report):
    path, _ = report
    dates = sheet(path, "Transactions")["Date"].astype(str).tolist()
    assert dates == sorted(dates)


def test_every_transaction_traces_back_to_its_source_file(report):
    path, _ = report
    tx = sheet(path, "Transactions")
    assert set(tx["Source File"]) == {f"{stem}.pdf" for stem in EXPECTED}
    assert tx["Source File"].notna().all()


# --- Summary --------------------------------------------------------------


def test_summary_has_one_row_per_statement_in_month_order(report):
    path, _ = report
    summary = sheet(path, "Summary")
    assert summary["Month"].tolist() == sorted(period for period, _, _ in EXPECTED.values())


def test_summary_totals_match_the_parsed_statements(report, parsed_by_stem):
    path, _ = report
    summary = sheet(path, "Summary").set_index("Month")

    for stem, (period, tx_count, net) in EXPECTED.items():
        row = summary.loc[period]
        assert row["Transactions"] == tx_count
        assert row["Net"] == pytest.approx(net, abs=0.01)
        assert row["Stated Net"] == pytest.approx(parsed_by_stem[stem].stated_net, abs=0.01)
        assert row["Reconciled"] == "yes"


def test_inflows_and_outflows_split_the_net(report):
    path, _ = report
    summary = sheet(path, "Summary")
    assert (summary["Inflows"] >= 0).all()
    assert (summary["Outflows"] <= 0).all()
    assert (summary["Inflows"] + summary["Outflows"]).round(2).tolist() == pytest.approx(
        summary["Net"].round(2).tolist(), abs=0.01
    )


# --- By Category ----------------------------------------------------------


def test_category_pivot_conserves_the_money(report, parsed_statements):
    """The pivot must not lose or double-count a single transaction."""
    path, _ = report
    by_category = sheet(path, "By Category")
    expected_total = round(
        sum(t["amount"] for st in parsed_statements for t in st.transactions), 2
    )
    assert by_category["Total"].sum() == pytest.approx(expected_total, abs=0.01)


def test_category_pivot_has_a_column_per_month(report):
    path, _ = report
    by_category = sheet(path, "By Category")
    months = {period for period, _, _ in EXPECTED.values()}
    assert months <= set(by_category.columns)
    assert by_category["Total"].tolist() == sorted(by_category["Total"].tolist())


# --- formatting -----------------------------------------------------------


def test_headers_are_bold_and_panes_frozen_on_every_sheet(report):
    path, _ = report
    workbook = openpyxl.load_workbook(path)
    for ws in workbook.worksheets:
        assert ws.freeze_panes == "A2"
        assert all(cell.font.bold for cell in ws[1])


def test_column_widths_are_set_but_capped(report):
    path, _ = report
    workbook = openpyxl.load_workbook(path)
    widths = [
        dim.width
        for ws in workbook.worksheets
        for dim in ws.column_dimensions.values()
        if dim.width is not None
    ]
    assert widths
    assert max(widths) <= 42


# --- not hardcoded to twelve months ---------------------------------------


def test_report_scales_down_to_a_subset(tmp_path, parsed_statements):
    subset = parsed_statements[:2]
    path = tmp_path / "subset.xlsx"
    summary = build_report(subset, path)

    assert len(summary) == 2
    tx = sheet(path, "Transactions")
    assert len(tx) == sum(len(st.transactions) for st in subset)
    assert set(tx["Month"]) == {st.period for st in subset}


# --- statements that yielded nothing --------------------------------------


def empty_statement(name="statement_2026_13.pdf", period=None) -> Statement:
    """What parse_statement returns for a PDF it could not read."""
    return Statement(
        source_file=name,
        period=period,
        transactions=[],
        stated_net=None,
        warnings=["no transactions extracted", "stated net movement not found"],
    )


def test_one_unreadable_statement_does_not_sink_the_report(tmp_path, parsed_statements):
    """It belongs in Summary/Notes, not as an exception that kills the run."""
    statements = [*parsed_statements[:2], empty_statement()]
    path = tmp_path / "with_empty.xlsx"
    summary = build_report(statements, path)

    assert len(summary) == 3
    tx = sheet(path, "Transactions")
    assert len(tx) == sum(len(st.transactions) for st in parsed_statements[:2])

    row = summary[summary["Transactions"] == 0].iloc[0]
    assert row["Reconciled"] == "NO"
    assert "no transactions extracted" in row["Notes"]
    assert row["Net"] == 0


def test_report_of_only_empty_statements_still_builds(tmp_path):
    path = tmp_path / "all_empty.xlsx"
    summary = build_report(
        [empty_statement("a.pdf", "2026-01"), empty_statement("b.pdf", "2026-02")], path
    )

    assert path.exists()
    assert list(summary.columns) == SUMMARY_COLUMNS
    assert summary["Transactions"].tolist() == [0, 0]
    assert summary["Month"].tolist() == ["2026-01", "2026-02"]


def test_empty_transactions_sheet_keeps_its_columns(tmp_path):
    path = tmp_path / "all_empty.xlsx"
    build_report([empty_statement()], path)

    tx = sheet(path, "Transactions")
    assert list(tx.columns) == TRANSACTION_COLUMNS
    assert tx.empty


def test_empty_report_keeps_all_three_sheets(tmp_path):
    path = tmp_path / "all_empty.xlsx"
    build_report([empty_statement()], path)

    workbook = openpyxl.load_workbook(path)
    assert workbook.sheetnames == ["Summary", "By Category", "Transactions"]
    for ws in workbook.worksheets:
        assert ws.freeze_panes == "A2"
        assert all(cell.font.bold for cell in ws[1])
